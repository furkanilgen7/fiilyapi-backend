"""TKL-B5.2 — teklif revizyonu Excel ciktisi: isveren / ic gorunum, SIZINTI BEKCISI.

Beklenen degerler API yanitindan (ve elle hesaplanmis sabitlerden) okunur; Excel uretim
fonksiyonundan turetilmez. Sizinti bekcisi tum hucreleri tarar (basliklar + degerler).
"""

from __future__ import annotations

import re
import uuid
import zipfile
from io import BytesIO

import openpyxl
import pytest

from app.core.access import AccessLevel, Scope
from app.modules.offers.export import (
    EMPLOYER_HEADERS,
    INTERNAL_HEADERS,
    ExportView,
    build_offer_workbook,
)
from app.modules.offers.models import Offer
from app.modules.offers.offer_read_schemas import OfferRevisionRead

from .._boq import _auth, _login_with_access, _set_permission
from ._offers import D, grup, kalem, rev_url, revizyon, teklif

#: Isveren ciktisinda HICBIR etiket/baslik hucresinde gecmemesi gereken kelimeler (kucuk harf).
YASAK_KELIMELER = ("maliyet", "genel gider", "kâr", "adam-saat", "referans", "son fiyat")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _export_url(offer_id: str, rev_no: int = 0) -> str:
    return rev_url(offer_id, rev_no) + "/export"


@pytest.fixture
async def dolu(client, admin, isveren, katalog) -> dict:
    """Iki gruplu Rev.0: [Beton 10 (maliyet 100), Demir 4 (maliyet 50, elle B.F. 80)] +
    [Kalip 2 (maliyetsiz → fiyatsiz)]. Taslak; kosullar dolu."""
    o = await teklif(client, admin, isveren, scope_summary="Kaba inşaat işleri")
    g1 = await grup(client, admin, o["id"], name="Kaba")
    g2 = await grup(client, admin, o["id"], name="İnce")
    await kalem(
        client, admin, o["id"], g1["id"], katalog[0].id, quantity="10", cost_unit_price="100"
    )
    await kalem(
        client,
        admin,
        o["id"],
        g1["id"],
        katalog[2].id,
        quantity="4",
        cost_unit_price="50",
        offer_unit_price="80",
    )
    await kalem(client, admin, o["id"], g2["id"], katalog[1].id, quantity="2")
    resp = await client.patch(
        rev_url(o["id"]),
        json={"payment_terms": "Peşin ödeme", "delivery_days": 90, "notes": "Nakliye dahil"},
        headers=admin,
    )
    assert resp.status_code == 200, resp.text
    return {"offer_id": o["id"], "offer_no": o["offer_no"]}


async def _indir(client, headers, offer_id: str, view: str | None = None, rev_no: int = 0):
    yol = _export_url(offer_id, rev_no) + (f"?view={view}" if view else "")
    return await client.get(yol, headers=headers)


def _kitap(resp) -> openpyxl.Workbook:
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith(XLSX)
    return openpyxl.load_workbook(BytesIO(resp.content))


def _hucreler(kitap: openpyxl.Workbook) -> list:
    return [c for satir in kitap.active.iter_rows() for c in satir if c.value is not None]


def _satirlar(kitap: openpyxl.Workbook) -> list[tuple]:
    return [tuple(c.value for c in satir) for satir in kitap.active.iter_rows()]


def _tablo(kitap: openpyxl.Workbook, basliklar: tuple[str, ...]) -> list[tuple]:
    """Baslik satirindan sonraki satirlar (genislik = baslik sayisi)."""
    satirlar = _satirlar(kitap)
    bas = next(i for i, s in enumerate(satirlar) if s[: len(basliklar)] == basliklar)
    return [s[: len(basliklar)] for s in satirlar[bas + 1 :]]


def _kalemler(rev: dict) -> list[dict]:
    return [k for g in rev["groups"] for k in g["items"]]


def _ic_degerler(rev: dict) -> set[str]:
    """Isverene GIZLI olmasi gereken sayisal degerlerin METIN temsilleri (API yanitindan)."""
    degerler: set[str] = set()
    for k in _kalemler(rev):
        for ad in ("cost", "overhead", "profit", "profit_pct", "man_hours"):
            if k["internal"][ad] is not None:
                degerler.add(k["internal"][ad])
        if k["cost_unit_price"] is not None:
            degerler.add(k["cost_unit_price"])
    for v in rev["totals"]["internal"].values():
        if v is not None:
            degerler.add(v)
    return degerler


def _musteri_degerler(rev: dict) -> set[str]:
    degerler = {k["quantity"] for k in _kalemler(rev) if k["quantity"] is not None}
    for k in _kalemler(rev):
        if k["customer"]:
            degerler |= {k["customer"]["unit_price"], k["customer"]["amount"]}
    degerler |= {v for v in rev["totals"]["customer"].values() if v is not None}
    return degerler


# --------------------------------------------------------------- SIZINTI BEKCISI


def _ham_parcalar(icerik: bytes) -> dict[str, str]:
    """xlsx = ZIP: TUM parcalar ham metin (sayfalar, sharedStrings, yorumlar, docProps, workbook…).
    Hucre degeri okuyan tarama gizli sayfa/yorum/ozellik/tanimli ad kacagini GOREMEZ."""
    with zipfile.ZipFile(BytesIO(icerik)) as zf:
        return {ad: zf.read(ad).decode("utf-8", errors="replace") for ad in zf.namelist()}


def _sayi_deseni(deger: str) -> re.Pattern[str]:
    """Rakam/nokta/virgulle BITISIK olmayan sayi: `1200.00` `11200.00`/`1200.001`/`1,200.00` icinde
    sayilmaz ama `(1200.00/...)` ya da `<t>1200.00</t>` icinde BULUNUR."""
    return re.compile(rf"(?<![\d.,]){re.escape(deger)}(?![\d])")


def _tr_kucuk(metin: str) -> str:
    return metin.replace("İ", "i").replace("I", "ı").lower()


def _yasak_kelimeler_bul(metin: str) -> list[str]:
    """Buyuk/kucuk harf ve Turkce İ/ı duyarli (`MALİYET`, `KÂR`, `Adam-Saat`…) + buyuk `GG`."""
    kucuk = {metin.lower(), _tr_kucuk(metin)}
    bulunan = [y for y in YASAK_KELIMELER if any(y in k for k in kucuk)]
    return bulunan + (["GG"] if "GG" in metin else [])


def _kitap_gizli_icerik_yok(kitap: openpyxl.Workbook) -> None:
    assert len(kitap.sheetnames) == 1, kitap.sheetnames  # gizli ikinci sayfa YOK
    assert not kitap.defined_names, list(kitap.defined_names)
    for sayfa in kitap.worksheets:
        assert sayfa.sheet_state == "visible"
        for satir in sayfa.iter_rows():
            for hucre in satir:
                assert hucre.comment is None, f"{hucre.coordinate}: yorum var"
        assert not [r for r, d in sayfa.row_dimensions.items() if d.hidden], "gizli satir"
        assert not [c for c, d in sayfa.column_dimensions.items() if d.hidden], "gizli sutun"
    p = kitap.properties
    for ad in ("description", "subject", "keywords", "title", "category", "comments"):
        assert not getattr(p, ad, None), f"belge ozelligi dolu: {ad}={getattr(p, ad)!r}"


async def test_isveren_xlsx_ic_degerleri_ve_ic_basliklari_HICBIR_hucrede_icermez(
    client, admin, dolu
) -> None:
    rev = await revizyon(client, admin, dolu["offer_id"])
    ic = _ic_degerler(rev)
    # on kosul: test verisi icin ic ve musteri degerleri CAKISMIYOR (aksi hâlde bekci anlamsiz)
    assert ic and not (ic & _musteri_degerler(rev)), ic & _musteri_degerler(rev)

    resp = await _indir(client, admin, dolu["offer_id"], "employer")
    kitap = _kitap(resp)
    hucreler = _hucreler(kitap)
    assert len(hucreler) > 30
    for hucre in hucreler:
        deger = str(hucre.value)
        assert deger not in ic, f"{hucre.coordinate}: ic deger SIZDI: {deger}"
        assert not _yasak_kelimeler_bul(deger), f"{hucre.coordinate}: {deger}"
    # (a) yapisal: gizli sayfa / yorum / belge ozelligi / tanimli ad / gizli satir-sutun YOK
    _kitap_gizli_icerik_yok(kitap)
    # (b) HAM tarama: TUM zip parcalari (hucre okuyan tarama bunlari gormez)
    parcalar = _ham_parcalar(resp.content)
    assert any(ad.startswith("xl/worksheets/") for ad in parcalar)
    assert {"xl/workbook.xml", "docProps/core.xml"} <= set(parcalar)  # sharedStrings olabilir
    for ad, metin in parcalar.items():
        if ad.startswith("xl/theme/"):  # yalniz renk/font tanimi; veri tasimaz
            continue
        for deger in ic:
            assert not _sayi_deseni(deger).search(metin), f"{ad}: ic deger SIZDI: {deger}"
        assert not _yasak_kelimeler_bul(metin), f"{ad}: {_yasak_kelimeler_bul(metin)}"
    # (c) varsayilan gorunum = isveren: ayni dosya
    varsayilan = _kitap(await _indir(client, admin, dolu["offer_id"]))
    assert _satirlar(varsayilan) == _satirlar(kitap)


async def test_ic_xlsx_ayni_degerleri_ICERIR_pozitif_kontrol(client, admin, dolu) -> None:
    rev = await revizyon(client, admin, dolu["offer_id"])
    ic = _ic_degerler(rev)
    resp = await _indir(client, admin, dolu["offer_id"], "internal")
    kitap = _kitap(resp)
    metinler = {str(h.value) for h in _hucreler(kitap)}
    assert ic <= metinler, ic - metinler
    # ic etiketler de var
    for etiket in ("Adam-saat", "Maliyet", "GG", "Kâr", "İÇ TOPLAMLAR", "Fiyatsız Kalem Sayısı"):
        assert etiket in metinler, etiket
    # ham tarama ayni degerleri BULUR (aksi hâlde isveren taramasi kor olabilirdi)
    ham = "\n".join(_ham_parcalar(resp.content).values())
    for deger in ic:
        assert _sayi_deseni(deger).search(ham), f"ham taramada bulunamadi: {deger}"
    assert {"maliyet", "kâr", "adam-saat", "GG"} <= set(_yasak_kelimeler_bul(ham))


def test_sayi_deseni_bitisik_rakam_nokta_virgul_saymaz() -> None:
    desen = _sayi_deseni("1200.00")
    assert desen.search("<t>1200.00</t>") and desen.search("Is (1200.00/340.00)")
    for degil in ("11200.00", "1200.001", "1,200.00", "0.1200.00"):
        assert not desen.search(degil), degil
    assert _yasak_kelimeler_bul("MALİYET") == ["maliyet"]
    assert _yasak_kelimeler_bul("Adam-Saat") == ["adam-saat"]
    assert _yasak_kelimeler_bul("SON FİYAT") == ["son fiyat"]
    assert _yasak_kelimeler_bul("Genel Gider") == ["genel gider"]
    assert _yasak_kelimeler_bul("KÂR") == ["kâr"]
    assert _yasak_kelimeler_bul("Poz No") == []


# --------------------------------------------------------------- icerik


async def test_basliklar_ve_sira_sabit(client, admin, dolu) -> None:
    emp = _kitap(await _indir(client, admin, dolu["offer_id"], "employer"))
    ic = _kitap(await _indir(client, admin, dolu["offer_id"], "internal"))
    # KAT-B3 / T49: "Bakanlık No" Poz No'nun HEMEN SAĞINDA (kullanıcı kararıyla eklendi)
    assert EMPLOYER_HEADERS == (
        "Poz No",
        "Bakanlık No",
        "İş Kalemi Tarifi",
        "Birim",
        "Miktar",
        "Teklif B.F.",
        "Tutar",
    )
    assert INTERNAL_HEADERS[7:] == (
        "Adam-saat",
        "Maliyet B.F.",
        "GG %",
        "Kâr %",
        "Maliyet",
        "GG",
        "Kâr",
    )
    assert any(s[: len(EMPLOYER_HEADERS)] == EMPLOYER_HEADERS for s in _satirlar(emp))
    assert any(s[: len(INTERNAL_HEADERS)] == INTERNAL_HEADERS for s in _satirlar(ic))
    assert not any(s[: len(INTERNAL_HEADERS)] == INTERNAL_HEADERS for s in _satirlar(emp))


FORMUL_KODU = '=HYPERLINK("http://k.co")'


@pytest.fixture
async def kodlu_katalog(seeded_db, katalog):
    """`dolu`dan ONCE istenmeli: kalem olusurken katalogdan `source_code` snapshot'lanir."""
    for entry, kod in zip(katalog, ["15.100.1001", None, FORMUL_KODU], strict=True):
        entry.source_code = kod
    await seeded_db.flush()
    return katalog


@pytest.mark.parametrize("gorunum", ["employer", "internal"])
async def test_bakanlik_no_sutunu_poz_no_yaninda_kodlu_deger_kodsuz_bos_formul_string(
    client, admin, kodlu_katalog, dolu, gorunum
) -> None:
    rev = await revizyon(client, admin, dolu["offer_id"])
    resp = await _indir(client, admin, dolu["offer_id"], gorunum)
    kitap = _kitap(resp)
    basliklar = EMPLOYER_HEADERS if gorunum == "employer" else INTERNAL_HEADERS
    assert basliklar[:2] == ("Poz No", "Bakanlık No")
    satirlar = {s[0]: s for s in _tablo(kitap, basliklar) if s[0] and s[2]}
    beklenen = {k["poz_no"]: k["source_code"] for k in _kalemler(rev)}
    assert sorted(v for v in beklenen.values() if v) == sorted(["15.100.1001", FORMUL_KODU])
    assert None in beklenen.values()  # kodsuz kalem de var
    for poz, kod in beklenen.items():
        assert satirlar[poz][1] == kod  # kodsuz → None (bos hucre)
    # formul degil STRING: veri tipi 's' (HYPERLINK calismaz)
    hucre = next(c for s in kitap.active.iter_rows() for c in s if c.value == FORMUL_KODU)
    assert hucre.data_type == "s"
    assert not any(c.data_type == "f" for s in kitap.active.iter_rows() for c in s)


async def test_kunye_toplamlar_ve_kosullar(client, admin, dolu) -> None:
    rev = await revizyon(client, admin, dolu["offer_id"])
    kitap = _kitap(await _indir(client, admin, dolu["offer_id"]))
    satirlar = _satirlar(kitap)
    kunye = {s[0]: s[1] for s in satirlar if s[0] and s[1] is not None}
    assert kunye["Teklif No"] == f"{dolu['offer_no']} Rev.0"
    assert kunye["İşveren"] == "Akın İnşaat A.Ş."
    assert kunye["Kapsam"] == "Kaba inşaat işleri"
    assert re.fullmatch(r"\d{2}\.\d{2}\.\d{4}", kunye["Tarih"])
    assert re.fullmatch(r"\d{2}\.\d{2}\.\d{4}", kunye["Geçerlilik Bitişi"])
    t = rev["totals"]["customer"]
    toplam = {s[0]: s[6] for s in satirlar if s[0] in ("NET (KDV Hariç)", "GENEL TOPLAM")}
    assert toplam == {"NET (KDV Hariç)": t["net"], "GENEL TOPLAM": t["gross"]}
    kdv = next(s for s in satirlar if s[0] and s[0].startswith("KDV (%"))
    assert kdv[0] == "KDV (%20)" and kdv[6] == t["vat"]
    assert kunye["Ödeme Koşulu"] == "Peşin ödeme"
    assert kunye["Teslim Süresi"] == "90 gün"
    assert kunye["Fiyat Farkı"] == "Endeksli (ÜFE)" or kunye["Fiyat Farkı"] == "Sabit fiyat"
    assert kunye["Notlar"] == "Nakliye dahil"
    # sabit sayilar (elle): net 1288.00 + 320.00 = 1608.00
    assert D(t["net"]) == D("1608.00")


async def test_tutarlar_API_ile_birebir_ve_her_hucre_str(client, admin, dolu) -> None:
    rev = await revizyon(client, admin, dolu["offer_id"])
    for gorunum in ("employer", "internal"):
        kitap = _kitap(await _indir(client, admin, dolu["offer_id"], gorunum))
        for hucre in _hucreler(kitap):
            assert isinstance(hucre.value, str), f"{hucre.coordinate}: {type(hucre.value)}"
            assert hucre.data_type == "s"
    kitap = _kitap(await _indir(client, admin, dolu["offer_id"], "internal"))
    satirlar = {s[0]: s for s in _tablo(kitap, INTERNAL_HEADERS) if s[0] and s[2]}
    for k in _kalemler(rev):
        satir = satirlar[k["poz_no"]]
        assert satir[4] == k["quantity"]
        assert satir[5] == (k["customer"]["unit_price"] if k["customer"] else None)
        assert satir[6] == (k["customer"]["amount"] if k["customer"] else None)
        assert satir[7] == k["internal"]["man_hours"]
        assert satir[8] == k["cost_unit_price"]
        assert (satir[11], satir[12], satir[13]) == (
            k["internal"]["cost"],
            k["internal"]["overhead"],
            k["internal"]["profit"],
        )


async def test_grup_ara_toplamlari_ve_fiyatsiz_kalem_bos_hucre(client, admin, dolu) -> None:
    rev = await revizyon(client, admin, dolu["offer_id"])
    g1, g2 = rev["groups"]
    kitap = _kitap(await _indir(client, admin, dolu["offer_id"], "employer"))
    ara = [s for s in _tablo(kitap, EMPLOYER_HEADERS) if s[0] == "Ara Toplam"]
    assert [s[6] for s in ara] == ["1608.00", "0.00"]  # ikinci grup yalniz fiyatsiz kalem
    fiyatsiz = next(k for k in g2["items"])
    assert fiyatsiz["priced"] is False
    satir = next(s for s in _tablo(kitap, EMPLOYER_HEADERS) if s[0] == fiyatsiz["poz_no"])
    assert satir[4] == fiyatsiz["quantity"] and satir[5] is None and satir[6] is None
    assert sum(D(k["customer"]["amount"]) for k in g1["items"]) == D(ara[0][6])
    # ic gorunum: ara toplam maliyet/GG/kar da dolu; fiyatsiz grup adam-saat DOLU
    ic = _kitap(await _indir(client, admin, dolu["offer_id"], "internal"))
    ara_ic = [s for s in _tablo(ic, INTERNAL_HEADERS) if s[0] == "Ara Toplam"]
    assert ara_ic[0][11] == "1200.00"  # 1000 + 200 (maliyet), elle
    assert ara_ic[1][7] == fiyatsiz["internal"]["man_hours"]


async def test_miktari_girilmemis_fiyatsiz_kalem_CALISMAZ_degil_BOS_basar(
    client, admin, dolu
) -> None:
    """B5.1 `quantity`yi NULL yapabilir: `priced=False` + `quantity=None` + `customer=None`."""
    rev = OfferRevisionRead.model_validate(await revizyon(client, admin, dolu["offer_id"]))
    g2 = rev.groups[1]
    bos = g2.items[0].model_copy(update={"quantity": None, "customer": None, "priced": False})
    grup2 = g2.model_copy(update={"items": [bos]})
    rev = rev.model_copy(update={"groups": [rev.groups[0], grup2]})
    offer = Offer(offer_no="TKL-2026-0001", employer_name="X", title="Y", scope_summary=None)
    for gorunum in ExportView:
        kitap = openpyxl.load_workbook(build_offer_workbook(offer, rev, gorunum))
        basliklar = EMPLOYER_HEADERS if gorunum is ExportView.employer else INTERNAL_HEADERS
        satir = next(s for s in _tablo(kitap, basliklar) if s[0] == bos.poz_no)
        assert satir[4:7] == (None, None, None)
        assert "None" not in {str(h.value) for h in _hucreler(kitap)}


# --------------------------------------------------------------- dosya adi


async def test_dosya_adi(client, admin, dolu) -> None:
    emp = await _indir(client, admin, dolu["offer_id"], "employer")
    ic = await _indir(client, admin, dolu["offer_id"], "internal")
    assert re.fullmatch(r"TKL-\d{4}-\d{4}", dolu["offer_no"])
    assert f"{dolu['offer_no']}-Rev0-isveren.xlsx" in emp.headers["content-disposition"]
    assert f"{dolu['offer_no']}-Rev0-ic.xlsx" in ic.headers["content-disposition"]
    assert emp.headers["content-disposition"].startswith("attachment;")


# --------------------------------------------------------------- hatalar / izin


async def test_olmayan_teklif_ve_revizyon_404_gecersiz_view_422(client, admin, dolu) -> None:
    assert (await _indir(client, admin, str(uuid.uuid4()))).status_code == 404
    assert (await _indir(client, admin, dolu["offer_id"], rev_no=7)).status_code == 404
    assert (await _indir(client, admin, dolu["offer_id"], "baska")).status_code == 422


async def test_kimliksiz_401(client, dolu) -> None:
    assert (await client.get(_export_url(dolu["offer_id"]))).status_code == 401


@pytest.mark.parametrize("role_key", ["site_chief", "field_engineer"])
async def test_contracts_yok_roller_403(client, db_session, user_factory, dolu, role_key) -> None:
    token = await _login_with_access(
        client, db_session, user_factory, role_key, f"{role_key}.{uuid.uuid4().hex[:6]}@tkl.co"
    )
    for gorunum in ("employer", "internal"):
        resp = await _indir(client, _auth(token), dolu["offer_id"], gorunum)
        assert resp.status_code == 403, gorunum


async def test_muhasebe_contracts_view_200_ve_tam_deger(
    client, admin, db_session, user_factory, dolu
) -> None:
    await _set_permission(db_session, "accounting", "contracts", AccessLevel.view, Scope.all)
    token = await _login_with_access(
        client, db_session, user_factory, "accounting", f"acc.{uuid.uuid4().hex[:6]}@tkl.co"
    )
    kitap = _kitap(await _indir(client, _auth(token), dolu["offer_id"], "internal"))
    assert "1608.00" in {str(h.value) for h in _hucreler(kitap)}


async def test_kisitli_disiplinli_kullanici_excel_indiremez_R5(
    client, db_session, user_factory, dolu
) -> None:
    """TKL-B4.6 (R5/SO-19): disiplin atanmis kullanici teklif Excel'ini (iki gorunum) indiremez;
    ayni rolun kisitsiz kullanicisi indirir (POZITIF KONTROL)."""
    from sqlalchemy import select

    from app.modules.catalog.models import ContractorType, EvDiscipline
    from app.modules.earned_value.models import UserDiscipline
    from app.modules.users.models import User

    disiplin = EvDiscipline(
        code="KSX", name="Kisitli", color="#2563EB", default_contractor_type=ContractorType.OWN
    )
    db_session.add(disiplin)
    await db_session.flush()
    eposta = f"pm.kisitli.{uuid.uuid4().hex[:6]}@tkl.co"
    token = await _login_with_access(client, db_session, user_factory, "project_manager", eposta)
    uid = (await db_session.execute(select(User.id).where(User.email == eposta))).scalar_one()
    db_session.add(UserDiscipline(user_id=uid, discipline_id=disiplin.id))
    await db_session.flush()
    for gorunum in ("employer", "internal"):
        resp = await _indir(client, _auth(token), dolu["offer_id"], gorunum)
        assert resp.status_code == 403, gorunum
    serbest = await _login_with_access(
        client, db_session, user_factory, "project_manager", f"pm.s.{uuid.uuid4().hex[:6]}@tkl.co"
    )
    for gorunum in ("employer", "internal"):
        resp = await _indir(client, _auth(serbest), dolu["offer_id"], gorunum)
        assert resp.status_code == 200, gorunum


# --------------------------------------------------------------- maske


async def test_limited_kapsamda_para_hucreleri_BOS_digerleri_gorunur(
    client, admin, db_session, user_factory, dolu
) -> None:
    rev = await revizyon(client, admin, dolu["offer_id"])
    para = (
        {k["customer"]["amount"] for k in _kalemler(rev) if k["customer"]}
        | {k["customer"]["unit_price"] for k in _kalemler(rev) if k["customer"]}
        | {v for v in rev["totals"]["customer"].values()}
        | _ic_degerler(rev)
        - {k["internal"]["man_hours"] for k in _kalemler(rev)}
        - {rev["totals"]["internal"]["man_hours"]}
    )
    await _set_permission(db_session, "accounting", "contracts", AccessLevel.view, Scope.limited)
    token = await _login_with_access(
        client, db_session, user_factory, "accounting", f"lim.{uuid.uuid4().hex[:6]}@tkl.co"
    )
    for gorunum in ("employer", "internal"):
        kitap = _kitap(await _indir(client, _auth(token), dolu["offer_id"], gorunum))
        metinler = {str(h.value) for h in _hucreler(kitap)}
        assert not (para & metinler), (gorunum, para & metinler)
        # miktar, poz, tarif, kosullar GORUNUR
        assert {k["quantity"] for k in _kalemler(rev)} <= metinler
        assert {k["poz_no"] for k in _kalemler(rev)} <= metinler
        assert "Peşin ödeme" in metinler
        satirlar = _satirlar(kitap)
        genel = next(s for s in satirlar if s[0] == "GENEL TOPLAM")
        assert genel[6] is None
        ara = next(s for s in satirlar if s[0] == "Ara Toplam")
        assert ara[6] is None  # kismi/maskeli toplam yazilmaz
    ic = _kitap(await _indir(client, _auth(token), dolu["offer_id"], "internal"))
    assert {k["internal"]["man_hours"] for k in _kalemler(rev)} <= {
        str(h.value) for h in _hucreler(ic)
    }


# --------------------------------------------------------------- ara toplam ↔ NET (V2)


@pytest.fixture
async def karisik(client, admin, isveren, katalog) -> dict:
    """Tek grup: [10 x maliyet 100 → tutar 1288.00] + [MIKTARSIZ, maliyet 50]. Miktarsiz kalem
    tutar uretmez: ara toplam NET ile (1288.00) tutarli olmali, BOSALMAMALI."""
    o = await teklif(client, admin, isveren)
    g = await grup(client, admin, o["id"], name="Karışık")
    await kalem(
        client, admin, o["id"], g["id"], katalog[0].id, quantity="10", cost_unit_price="100"
    )
    await kalem(client, admin, o["id"], g["id"], katalog[2].id, quantity=None, cost_unit_price="50")
    return o


async def test_miktarsiz_kalem_ara_toplami_BOSALTMAZ_NET_ile_tutarli(
    client, admin, karisik
) -> None:
    rev = await revizyon(client, admin, karisik["id"])
    assert rev["totals"]["unquantified_count"] == 1
    net = rev["totals"]["customer"]["net"]
    assert D(net) == D("1288.00")  # elle: 100 x 1,12 x 1,15 = 128,80; x 10
    emp = _kitap(await _indir(client, admin, karisik["id"], "employer"))
    ara = [s for s in _tablo(emp, EMPLOYER_HEADERS) if s[0] == "Ara Toplam"]
    assert [s[6] for s in ara] == ["1288.00"] and ara[0][6] == net
    miktarsiz = next(k for k in _kalemler(rev) if k["quantity"] is None)
    satir = next(s for s in _tablo(emp, EMPLOYER_HEADERS) if s[0] == miktarsiz["poz_no"])
    assert satir[4] is None and satir[6] is None  # kalem hucreleri yine bos
    # ic gorunum: maliyet/GG/kar ara toplamlari da dolu (elle 1000 / 120 / 168), adam-saat yalniz
    # miktarli kalemden (10 x 1,5 = 15); miktarsiz kalemin adam-saati bos
    ic = _kitap(await _indir(client, admin, karisik["id"], "internal"))
    ara_ic = next(s for s in _tablo(ic, INTERNAL_HEADERS) if s[0] == "Ara Toplam")
    assert (D(ara_ic[11]), D(ara_ic[12]), D(ara_ic[13])) == (D("1000"), D("120"), D("168"))
    assert D(ara_ic[7]) == D("15") and D(ara_ic[7]) == D(rev["totals"]["internal"]["man_hours"])
    satir_ic = next(s for s in _tablo(ic, INTERNAL_HEADERS) if s[0] == miktarsiz["poz_no"])
    assert satir_ic[7] is None  # adam-saat bilinmiyor: bos hucre (0 DEGIL)


async def test_limited_kapsamda_karisik_grubun_ara_toplami_BOS_adam_saat_gorunur(
    client, admin, db_session, user_factory, karisik
) -> None:
    """Gercekten maskeli (miktarli+fiyatli kalemin tutari gizli) → ara toplam BOS; kimlik kovasi
    olan adam-saat toplami maskelenmez."""
    await _set_permission(db_session, "accounting", "contracts", AccessLevel.view, Scope.limited)
    token = await _login_with_access(
        client, db_session, user_factory, "accounting", f"kar.{uuid.uuid4().hex[:6]}@tkl.co"
    )
    for gorunum, basliklar in (("employer", EMPLOYER_HEADERS), ("internal", INTERNAL_HEADERS)):
        kitap = _kitap(await _indir(client, _auth(token), karisik["id"], gorunum))
        ara = next(s for s in _tablo(kitap, basliklar) if s[0] == "Ara Toplam")
        assert ara[6] is None, gorunum
        if gorunum == "internal":
            assert ara[11] is None and ara[12] is None and ara[13] is None
            assert D(ara[7]) == D("15")
