"""EV-BORC-7 — QURR Excel'i JSON agaciyla AYNI sirayi izler + toplam satirinda kod hucresi.

Kurgu: conftest `baseline` + `saha_gunu` (KAB: Betonarme/I1 · DUV: Duvar/I2) — iki disiplin,
iki grup; duz yazim (once tum kalemler, sonra tum toplamlar) bu kurguda agac sirasindan AYRILIR.
Beklenen sira JSON'dan `parent_id` ile BAGIMSIZ kurulur (uretim siralayicisi kullanilmaz).
"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.modules.earned_value import report_qurr
from app.modules.earned_value.schemas_reports import QurrReport, QurrRow, QurrTotal

from .test_evborc3_s315 import _get, bugun  # noqa: F401 - fixture
from .test_reports import _allocate, _rep

_HEADER_ROW = 2


def _expected_tree(q: dict) -> list[tuple[str | None, str]]:
    """JSON → agac sirasi: disiplin { grup { kalemler, grup toplami }, disiplin toplami },
    sonra Σ D, Σ D+DL. Satir kimligi (kod, ad)."""
    rows, totals = q["rows"], q["totals"]
    out: list[tuple[str | None, str]] = []
    for d in (t for t in totals if t["kind"] == "discipline"):
        for g in (t for t in totals if t["kind"] == "group" and t["parent_id"] == d["node_id"]):
            out += [(r["code"], r["name"]) for r in rows if r["parent_id"] == g["node_id"]]
            out.append((g["code"], g["name"]))
        out.append((d["code"], d["name"]))
    out += [(t["code"], t["name"]) for t in totals if t["kind"] in ("direct_total", "all_total")]
    return out


async def _sheet(client, headers, santiye) -> list[tuple[str | None, str]]:  # noqa: ANN001
    resp = await client.get(_rep(santiye, "/weekly.xlsx"), headers=headers, params={"week": 1})
    assert resp.status_code == 200, resp.text
    ws = load_workbook(BytesIO(resp.content)).active
    header = [c.value for c in ws[_HEADER_ROW]]
    kod, ad = header.index("Kod"), header.index("İş tipi")
    return [(r[kod], r[ad]) for r in ws.iter_rows(min_row=_HEADER_ROW + 1, values_only=True)]


async def test_xlsx_rows_follow_json_tree_order(
    client,
    saha,
    admin,
    santiye,
    boq,
    baseline,
    saha_gunu,
    bugun,  # noqa: F811
) -> None:
    await _allocate(client, saha, santiye, boq, saha_gunu)
    q = await _get(client, admin, _rep(santiye, "/weekly"), week=1)
    expected = _expected_tree(q)
    assert len(expected) == len(q["rows"]) + len(q["totals"]), "kurgu: her satir agacta"
    assert len({d for d in (t["parent_id"] for t in q["totals"]) if d}) >= 2, "kurgu: 2 disiplin"
    names = [name for _, name in expected]
    assert len(names) == len(set(names)), "kurgu varsayimi: sayfadaki TUM adlar tekil"
    sheet = await _sheet(client, admin, santiye)
    assert sheet == expected  # (kod, ad) demeti — yalniz ad degil


async def test_xlsx_total_rows_carry_code(
    client,
    saha,
    admin,
    santiye,
    boq,
    baseline,
    saha_gunu,
    bugun,  # noqa: F811
) -> None:
    await _allocate(client, saha, santiye, boq, saha_gunu)
    q = await _get(client, admin, _rep(santiye, "/weekly"), week=1)
    rows = await _sheet(client, admin, santiye)
    names = [name for _, name in rows]
    assert len(names) == len(set(names)), (
        "kurgu varsayimi: sayfadaki TUM adlar tekil (ad ile eslesiyor)"
    )
    sheet = {name: code for code, name in rows}
    for t in q["totals"]:
        assert sheet[t["name"]] == t["code"], (t["kind"], t["name"])
    assert {t["code"] for t in q["totals"] if t["kind"] == "group"} == {"1", "2"}


# --- tree_lines birim testleri (EV-BORC-7 (1): eslenmeyen kalem bekcisi) -----------------


def _min_row(node_id: str, parent_id: str, code: str = "1.001", name: str = "Kalem") -> QurrRow:
    return QurrRow.model_construct(
        node_id=node_id, level=3, parent_id=parent_id, code=code, name=name
    )


def _min_total(
    kind: str, node_id: str | None, name: str, parent_id: str | None = None
) -> QurrTotal:
    return QurrTotal.model_construct(kind=kind, node_id=node_id, name=name, parent_id=parent_id)


def test_tree_lines_raises_when_row_parent_has_no_group_total() -> None:
    """Bir kalemin `parent_id`'si hicbir grup toplaminin `node_id`'sine
    eslenmiyorsa `tree_lines` ValueError firlatir (Excel sessizce kalemi atlayamaz)."""
    row = _min_row("i:eksik", parent_id="g:yok")
    report = QurrReport.model_construct(rows=[row], totals=[])
    with pytest.raises(ValueError, match="grup toplami olmayan"):
        report_qurr.tree_lines(report)


def test_tree_lines_order_is_item_group_discipline_direct_all() -> None:
    """Eslenen durumda sira: kalem, grup toplami, disiplin toplami, Σ D, Σ D+DL."""
    row = _min_row("i:1", parent_id="g:1", code="1.001", name="Kalem")
    group = _min_total("group", "g:1", "Grup", parent_id="d:1")
    discipline = _min_total("discipline", "d:1", "Disiplin", parent_id=None)
    direct_total = _min_total("direct_total", None, "Σ Doğrudan")
    all_total = _min_total("all_total", None, "Σ Doğrudan + Dolaylı")
    report = QurrReport.model_construct(
        rows=[row], totals=[group, discipline, direct_total, all_total]
    )
    lines = report_qurr.tree_lines(report)
    assert [(getattr(line, "code", None), line.name) for line in lines] == [
        ("1.001", "Kalem"),
        (None, "Grup"),
        (None, "Disiplin"),
        (None, "Σ Doğrudan"),
        (None, "Σ Doğrudan + Dolaylı"),
    ]


# --- 409 beyani (EV-BORC-7 (2)) ---------------------------------------------------------


async def test_approve_daily_without_diary_is_real_409(
    client,
    admin,
    santiye,
    boq,
    baseline,
) -> None:
    """Onay ucunun 'gunluk yok' dali HTTP duzeyinde bekcili: baseline'li ama o gun gunlugu
    HIC olmayan bir günde POST .../approve → 409, govde `report_daily.NOT_GENERATED` ile
    BIREBIR (kayma yok)."""
    from app.modules.earned_value import report_daily

    resp = await client.post(_rep(santiye, "/daily/2026-05-06/approve"), headers=admin)
    assert resp.status_code == 409, resp.text
    assert resp.json()["detail"] == report_daily.NOT_GENERATED


def test_ev_report_routes_declare_their_real_409() -> None:
    """Ucun GERCEKTEN dondugu 409 openapi'de beyanli; metin mesaj sabitiyle AYNI (kayma yok).
    Gercek donus bekcileri: test_S6_no_baseline_is_409_not_404 · onay 409 testleri."""
    from app.main import app
    from app.modules.earned_value.diary_adapter import NO_BASELINE
    from app.modules.earned_value.guards import SITE_COMPLETED_BUDGET_READ_ONLY
    from app.modules.earned_value.report_daily import NOT_GENERATED

    paths = app.openapi()["paths"]
    base = "/sites/{site_id}/earned-value/reports"
    beklenen = {
        (f"{base}/weekly", "get"): (NO_BASELINE,),
        (f"{base}/weekly.xlsx", "get"): (NO_BASELINE,),
        (f"{base}/daily/{{day}}/approve", "post"): (
            NOT_GENERATED,
            SITE_COMPLETED_BUDGET_READ_ONLY,
        ),
    }
    for (path, method), metinler in beklenen.items():
        responses = paths[path][method]["responses"]
        assert "409" in responses, f"{method.upper()} {path}: 409 beyansiz"
        for metin in metinler:
            assert metin in responses["409"]["description"], (path, metin)
