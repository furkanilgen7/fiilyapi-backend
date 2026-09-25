"""EV-BORC-3 ek — spec §3.15 (frontend F3 planı): S1 S2 S4 S5 S6 S7 S9 S22 + XLSX sayı hücresi.

Kurgu: conftest `baseline` + `saha_gunu` (05.05.2026; KAB kendi: Betonarme/I1 · DUV taşeron:
Duvar/I2). Hepsi yalnız EK alan; S6 `week` OPSİYONEL (verilirse davranış aynı).
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.modules.earned_value import diary_adapter

from .conftest import DAY
from .test_reports import _allocate, _rep

D = Decimal


async def _get(client, headers, url: str, **params) -> dict:  # noqa: ANN001, ANN003
    resp = await client.get(url, headers=headers, params=params)
    assert resp.status_code == 200, resp.text
    return resp.json()


@pytest.fixture
def bugun(monkeypatch):  # noqa: ANN001, ANN201
    """QURR "bugün"ü (as_of + S6 varsayılan hafta) sabitlenir: hafta 1 içinde 06.05."""
    from app.core import timezone
    from app.modules.earned_value import report_qurr

    fixed = date(2026, 5, 6)
    monkeypatch.setattr(timezone, "today", lambda: fixed)
    monkeypatch.setattr(report_qurr, "today", lambda: fixed)
    return fixed


async def test_S1_S2_qurr_tree_parents_codes_mix_and_bands(
    client, saha, admin, santiye, boq, baseline, saha_gunu, disiplinler, bugun
) -> None:
    await _allocate(client, saha, santiye, boq, saha_gunu)
    q = await _get(client, admin, _rep(santiye, "/weekly"), week=1)
    g1, g2 = f"g:{boq['g1'].id}", f"g:{boq['g2'].id}"
    kab, duv = f"d:{disiplinler[0].id}", f"d:{disiplinler[1].id}"
    assert {r["node_id"]: r["parent_id"] for r in q["rows"]} == {
        f"i:{boq['i1'].id}": g1,
        f"i:{boq['i2'].id}": g2,
    }
    by = {(t["kind"], t["node_id"]): t for t in q["totals"]}
    assert (by[("group", g1)]["parent_id"], by[("group", g2)]["parent_id"]) == (kab, duv)
    assert by[("discipline", kab)]["parent_id"] is None
    assert (by[("group", g1)]["code"], by[("discipline", kab)]["code"]) == ("1", "KAB")
    assert (by[("group", g1)]["contractor_mix"], by[("discipline", duv)]["contractor_mix"]) == (
        "own",
        "subcon",
    )
    assert by[("direct_total", None)]["contractor_mix"] is None
    # PF 10/17 = 0,588 < 0,95 → kırmızı; taşeron (harcama yok) → bant yok
    assert (by[("group", g1)]["q_band"], by[("group", g1)]["r_band"]) == ("red", "red")
    assert by[("group", g2)]["q_band"] is None
    assert q["has_field_data"] is True  # S7


async def test_S4_quantity_tree_carries_pf_cum(
    client, saha, admin, santiye, boq, baseline, saha_gunu
) -> None:
    await _allocate(client, saha, santiye, boq, saha_gunu)
    rep = await _get(client, admin, _rep(santiye, "/daily"), date=DAY.isoformat())
    item = next(r for r in rep["quantities"] if r["node_id"] == f"i:{boq['i1'].id}")
    assert item["pf_cum"] is not None, "miktar ağacında pf_cum yok"
    assert (D(item["pf_cum"]), item["pf_cum_band"]) == (D(10) / D(17), "red")


async def test_S5_S22_histogram_week_and_curve_status(
    client, saha, admin, santiye, boq, baseline, saha_gunu
) -> None:
    await _allocate(client, saha, santiye, boq, saha_gunu)
    p = await _get(client, admin, f"/sites/{santiye.id}/earned-value/panel", date=DAY.isoformat())
    assert [(w["week_no"], w["week_start"], w["week_end"]) for w in p["histogram"]][:2] == [
        (1, "2026-05-04", "2026-05-10"),
        (2, "2026-05-11", "2026-05-17"),
    ]
    today = next(c for c in p["s_curve"] if c["day"] == DAY.isoformat())
    variance = D(today["progress_pct_cum"]) - D(today["planned_pct_cum"])
    assert today["variance"] is not None, "S-eğrisi noktasında sapma yok"
    assert D(today["variance"]) == variance and today["status"] == "late"
    future = next(c for c in p["s_curve"] if c["is_future"])
    assert (future["variance"], future["status"]) == (None, None)


async def test_S6_week_optional_defaults_to_today_and_errors_are_distinct(
    client, admin, santiye, boq, baseline, saha_gunu, bugun
) -> None:
    implicit = await _get(client, admin, _rep(santiye, "/weekly"))
    explicit = await _get(client, admin, _rep(santiye, "/weekly"), week=1)
    assert implicit["week_no"] == 1  # 06.05 → hafta 1
    for key in ("generated_at",):
        implicit.pop(key), explicit.pop(key)
    assert implicit == explicit  # week verilirse davranış AYNI
    missing = await client.get(_rep(santiye, "/weekly"), headers=admin, params={"week": 99})
    assert missing.status_code == 404


async def test_S6_no_baseline_is_409_not_404(client, admin, santiye, boq) -> None:
    resp = await client.get(_rep(santiye, "/weekly"), headers=admin)
    assert resp.status_code == 409, resp.text
    assert resp.json()["detail"] == diary_adapter.NO_BASELINE


async def test_S7_has_field_data_false_without_entries(
    client, admin, santiye, boq, baseline, bugun
) -> None:
    q = await _get(client, admin, _rep(santiye, "/weekly"), week=1)
    assert q["has_field_data"] is False


async def test_S9_overrun_warning_has_qty_cum_and_planned(
    client, admin, seeded_db, santiye, boq, baseline, saha_gunu
) -> None:
    saha_gunu["line"].quantity = D(15)
    await seeded_db.flush()
    rep = await _get(client, admin, _rep(santiye, "/daily"), date=DAY.isoformat())
    (w,) = [x for x in rep["warnings"] if x["code"] == "qty_overrun"]
    assert w.get("qty_cum") is not None and w.get("planned_qty") is not None, w
    assert (D(w["qty_cum"]), D(w["planned_qty"]), D(w["value"])) == (15, 10, 5)


async def test_xlsx_cells_are_numbers_with_formats(
    client, saha, admin, santiye, boq, baseline, saha_gunu, bugun
) -> None:
    await _allocate(client, saha, santiye, boq, saha_gunu)
    resp = await client.get(_rep(santiye, "/weekly.xlsx"), headers=admin, params={"week": 1})
    assert resp.status_code == 200
    ws = load_workbook(BytesIO(resp.content)).active
    header = [c.value for c in ws[2]]
    row = {h: c for h, c in zip(header, ws[3], strict=True)}  # I1 satırı
    assert row["İş tipi"].value == "Beton" and row["İş tipi"].data_type == "s"
    assert row["g Bütçe a-s"].data_type == "n" and row["g Bütçe a-s"].value == 200
    assert row["g Bütçe a-s"].number_format == "#,##0.00"
    assert row["c Gerçekleşen"].number_format == "#,##0.000"
    assert row["q PF"].number_format == "0.00" and abs(row["q PF"].value - 10 / 17) < 1e-9
    assert row["a Önceki miktar"].value is None  # önceki revizyon yok → BOŞ hücre (""" değil)


def test_S1_frozen_tree_orders_groups_by_boq_sort_not_uuid() -> None:
    """🔴 Bulgu: donmuş ağaç grupları UUID METNİNE göre sıralıyordu → dondurulmuş görünümde
    grup sırası ve konum kodu (GroupOut.code, QURR totals.code) RASTGELE. Kimlik sırası
    BOQ sırasının TERSİ kurulur; sıra BOQ `sort_order`ını izlemeli."""
    import uuid
    from types import SimpleNamespace

    from app.modules.earned_value.budget_snapshot import _snapshot_boq

    first, second = uuid.UUID(int=2**128 - 1), uuid.UUID(int=1)  # UUID sırası: second < first

    def leaf(group, code):  # noqa: ANN001, ANN202
        return SimpleNamespace(
            boq_group_id=group,
            boq_item_id=uuid.uuid4(),
            item_code=code,
            item_description=code,
            uom="m3",
            planned_qty=D(1),
            section_id=None,
            section_name=None,
        )

    boq = _snapshot_boq(
        [leaf(first, "01.001"), leaf(second, "02.001")],
        {first: ("Betonarme", 1), second: ("Duvar", 2)},
    )
    assert [(g.id, g.name) for g in boq.groups] == [(first, "Betonarme"), (second, "Duvar")]
