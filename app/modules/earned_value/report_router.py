"""Rapor uclari — `/sites/{site_id}/earned-value/reports/…` + panel (PLN-B3).

Kapilar: okuma VIEW · rapor ONAYI APPROVE (§2; onay o tarihe kadar gunluk + puantaji
kilitler — B2 kilit tablosuna yazar). Tamamlanmis santiyede onay 409 (B3.0 tek yardimci).
PDF sunucuda YOK (CEO: A4 yatay yazdirma gorunumu frontend'de); QURR Excel backend'de
(openpyxl prod kilidinde VAR).
"""

from __future__ import annotations

import uuid
from datetime import date
from io import BytesIO
from typing import Annotated

from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import Response
from openpyxl import Workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.db import get_db
from app.core.deps import get_current_user
from app.core.errors import ConflictError, NotFoundError
from app.core.http import content_disposition
from app.core.openapi import COMMON_ERROR_RESPONSES
from app.core.ratelimit import client_ip
from app.modules.audit.models import AuditAction
from app.modules.audit.service import record_audit
from app.modules.earned_value import audit_messages as msg
from app.modules.earned_value import report_daily, report_panel, report_qurr
from app.modules.earned_value.access import (
    APPROVE,
    VIEW,
    SiteContext,
    completed_site_guard,
    visible_site,
)
from app.modules.earned_value.diary_adapter import NO_BASELINE
from app.modules.earned_value.engine import ContractorType
from app.modules.earned_value.ev_input import build_site_input
from app.modules.earned_value.guards import SITE_COMPLETED_BUDGET_READ_ONLY
from app.modules.earned_value.schemas_reports import (
    ApprovalResult,
    DailyReport,
    PanelReport,
    QurrReport,
)
from app.modules.users.models import User

router = APIRouter(tags=["earned-value"], responses=COMMON_ERROR_RESPONSES)

_User = Annotated[User, Depends(get_current_user)]
_Db = Annotated[AsyncSession, Depends(get_db)]
_Writable = Annotated[SiteContext, Depends(completed_site_guard(SITE_COMPLETED_BUDGET_READ_ONLY))]
_BASE = "/sites/{site_id}/earned-value/reports"
NO_WEEK = "Hafta proje takviminde yok"
#: §3.15 S6: baseline yoksa NO_WEEK DEGIL, `diary_adapter.NO_BASELINE` (409, durum) — tek metin.
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@router.get("/sites/{site_id}/earned-value/panel", response_model=PanelReport, dependencies=[VIEW])
async def get_panel(
    site_id: uuid.UUID,
    user: _User,
    session: _Db,
    day: Annotated[date, Query(alias="date")],
    range_: Annotated[report_panel.Range, Query(alias="range")] = "4w",
    discipline_id: Annotated[str | None, Query(max_length=80)] = None,
    contractor_type: ContractorType | None = None,
) -> PanelReport:
    """Planlama paneli — filtre (disiplin kokü `d:…`, kendi/taseron) BUTUN panele uygulanir."""
    await visible_site(session, user, site_id)
    query = report_panel.PanelQuery(day, range_, discipline_id, contractor_type)
    return await report_panel.build_panel(session, site_id, query)


@router.get(f"{_BASE}/daily", response_model=DailyReport, dependencies=[VIEW])
async def get_daily_report(
    site_id: uuid.UUID, user: _User, session: _Db, day: Annotated[date, Query(alias="date")]
) -> DailyReport:
    """Gunluk ilerleme raporu (GIR). Onayli + kilitli gun → donmus snapshot (B3-5)."""
    await visible_site(session, user, site_id)
    return await report_daily.build_daily(session, site_id, day)


@router.post(
    f"{_BASE}/daily/{{day}}/approve", response_model=ApprovalResult, dependencies=[APPROVE]
)
async def approve_daily_report(
    request: Request, site_id: uuid.UUID, day: date, ctx: _Writable, user: _User, session: _Db
) -> ApprovalResult:
    """Onayla ve kilitle: K13 taslak gunluk varsa 422 · gunluk yoksa 409 · eksik gunler
    (gunlugu hic olmayan is gunleri) yanitta `missing_diary_dates` (B3-1)."""
    result = await report_daily.approve(session, site_id, day, user)
    await record_audit(
        session,
        action=AuditAction.approve,
        detail=msg.report_approved(
            ctx.project.name, ctx.site.name, day, result.report.version or 1
        ),
        actor_user_id=user.id,
        ip_address=client_ip(request),
    )
    return result


async def _qurr(session: AsyncSession, site_id: uuid.UUID, week: int | None) -> QurrReport:
    from app.core.timezone import today

    site = await build_site_input(session, site_id, today())
    if site is None:
        raise ConflictError(NO_BASELINE)
    report = await report_qurr.build_qurr(session, site_id, site, week)
    if report is None:
        raise NotFoundError(NO_WEEK)
    return report


@router.get(f"{_BASE}/weekly", response_model=QurrReport, dependencies=[VIEW])
async def get_weekly_report(
    site_id: uuid.UUID,
    user: _User,
    session: _Db,
    week: Annotated[int | None, Query(ge=1, le=600)] = None,
) -> QurrReport:
    """Haftalik QURR (B3-2: onaylanmaz, canli). `week` yoksa bugunun haftasi (takvime
    kirpilir, §3.15 S6). Baseline yok → 409 NO_BASELINE · hafta takvimde yok → 404 NO_WEEK."""
    await visible_site(session, user, site_id)
    return await _qurr(session, site_id, week)


#: Hucre bicimleri (§3.15: hucreler SAYI — Excel'de toplanabilir; gorunum bicimle).
_QTY, _MHR, _RATE, _PF = "#,##0.000", "#,##0.00", "0.0000", "0.00"
#: (baslik, alan, sayi bicimi | None = metin)
_COLUMNS = (
    ("Kod", "code", None),
    ("İş tipi", "name", None),
    ("Birim", "uom", None),
    ("a Önceki miktar", "a_prev_qty", _QTY),
    ("b Miktar", "b_qty", _QTY),
    ("c Gerçekleşen", "c_qty_cum", _QTY),
    ("d Kalan", "d_remaining_qty", _QTY),
    ("e Bu hafta", "e_qty_week", _QTY),
    ("f Önceki a-s", "f_prev_budget_mhr", _MHR),
    ("g Bütçe a-s", "g_budget_mhr", _MHR),
    ("h Kazanılan", "h_earned_cum", _MHR),
    ("i Harcanan", "i_spent_cum", _MHR),
    ("j Kalan a-s", "j_remaining_mhr", _MHR),
    ("k Bu hafta kaz.", "k_earned_week", _MHR),
    ("l Bu hafta harc.", "l_spent_week", _MHR),
    ("m Önceki oran", "m_prev_unit_mhr", _RATE),
    ("n Oran", "n_unit_mhr", _RATE),
    ("o Gerç. oran", "o_actual_unit_mhr_cum", _RATE),
    ("p Hafta oranı", "p_actual_unit_mhr_week", _RATE),
    ("q PF", "q_pf_cum", _PF),
    ("r Hafta PF", "r_pf_week", _PF),
)
#: Toplam satirinda dolu olan alanlar (f–l + q, r); digerleri bos hucre.
_TOTAL_FIELDS = frozenset(
    {
        "f_prev_budget_mhr",
        "g_budget_mhr",
        "h_earned_cum",
        "i_spent_cum",
        "j_remaining_mhr",
        "k_earned_week",
        "l_spent_week",
        "q_pf_cum",
        "r_pf_week",
    }
)


def _append(ws, values: list) -> None:  # noqa: ANN001
    """Satiri yazar; sayisal hucreye kolonun bicimini uygular (None → bos hucre)."""
    ws.append(values)
    for cell, (_, _, fmt) in zip(ws[ws.max_row], _COLUMNS, strict=True):
        if fmt is not None and cell.value is not None:
            cell.number_format = fmt


def qurr_workbook(report: QurrReport) -> bytes:
    """Saf sunum: hesap YAPMAZ, degerleri yazar. Sayilar SAYI hucresidir (Decimal; §3.15)."""
    wb = Workbook()
    ws = wb.active
    ws.title = f"QURR H{report.week_no}"
    ws.append([f"Haftalık QURR · H{report.week_no} · {report.week_start} – {report.week_end}"])
    ws.append([label for label, _, _ in _COLUMNS])
    for row in report.rows:
        _append(ws, [getattr(row, attr) for _, attr, _ in _COLUMNS])
    for t in report.totals:
        _append(
            ws,
            [
                t.name if attr == "name" else getattr(t, attr) if attr in _TOTAL_FIELDS else None
                for _, attr, _ in _COLUMNS
            ],
        )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get(
    f"{_BASE}/weekly.xlsx",
    dependencies=[VIEW],
    response_class=Response,
    responses={200: {"content": {XLSX: {}}}},
)
async def export_weekly_report(
    site_id: uuid.UUID,
    user: _User,
    session: _Db,
    week: Annotated[int | None, Query(ge=1, le=600)] = None,
) -> Response:
    """QURR Excel — okuma ucuyla AYNI hesap (`_qurr`); export saf sunumdur."""
    ctx = await visible_site(session, user, site_id)
    report = await _qurr(session, site_id, week)
    return Response(
        content=qurr_workbook(report),
        media_type=XLSX,
        headers={
            "Content-Disposition": content_disposition(
                f"QURR-{ctx.site.name}-H{report.week_no}.xlsx"
            )
        },
    )
