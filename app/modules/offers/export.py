"""Teklif revizyonu Excel ciktisi — IKI tur (TKL-B5.2, T36-T37).

Saf sunum katmani: DB/`Request` bilmez; yalniz `OfferRevisionRead` (teklif okuma yolunun
urettigi, kapsam maskesinden gecmis zarf) calisma kitabina cevirir. Ikinci bir hesap YOKTUR:
her deger zarftan AYNEN okunur; tek turetme grup ara toplamidir (zarftaki kalem tutarlarinin
toplami — tutarlar 0,01'e yuvarli oldugu icin tam esittir).

## Iki tur
* `employer` (ISVEREN): Poz No · Bakanlık No · Tarif · Birim · Miktar · Teklif B.F. · Tutar +
  net/KDV/genel toplam + kosullar. Maliyet, GG, kar, kar %, adam-saat, katalog/referans/son
  fiyat YOKTUR.
* `internal` (IC): isveren sutunlarina ek Adam-saat · Maliyet B.F. · GG % · Kar % · Maliyet ·
  GG · Kar + ic toplamlar.

🔴 SIZINTI SINIRI: isveren sayfasi `item.customer` / `totals.customer` ALT NESNELERINDEN ve
yapisal alanlardan (poz, tarif, birim, miktar) okur; `internal`, `cost_unit_price`,
`overhead_pct`, `profit_pct` isveren yolunda HIC OKUNMAZ. Bekcisi
`tests/modules/offers/test_offer_export.py` (her hucre taranir).

FLOAT-YASAK (boq/export.py emsali): her hucre ACIKCA `str`; maskeli/girilmemis deger (`None`)
hucreye YAZILMAZ — hucre BOS kalir. `str(None)` ya da "0" yazmak sessizce yanlistir.
"""

from __future__ import annotations

from collections.abc import Callable, Sequence
from decimal import Decimal
from enum import Enum
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app.core.xlsx_text import write_text_cell
from app.modules.offers.models import Offer, OfferPriceEscalation
from app.modules.offers.offer_read_schemas import OfferItemRead, OfferRevisionRead
from app.modules.projects.models import PriceIndexType


class ExportView(str, Enum):
    employer = "employer"
    internal = "internal"


FILE_SUFFIX = {ExportView.employer: "isveren", ExportView.internal: "ic"}

SHEET_TITLE = "Teklif"
DATE_FORMAT = "%d.%m.%Y"

#: Isveren tablosu. Sira ve metin degistirilmez; ISTISNA: "Bakanlık No" Poz No'nun hemen sagina
#: KULLANICI KARARIYLA eklendi (KAT-B3 / T49-Q1; deger = kalemin `source_code` snapshot'i).
#: Sutun indeksleri (`_AMOUNT_COL`, ara toplam, genislik) BU listeden turetilir.
EMPLOYER_HEADERS: tuple[str, ...] = (
    "Poz No",
    "Bakanlık No",
    "İş Kalemi Tarifi",
    "Birim",
    "Miktar",
    "Teklif B.F.",
    "Tutar",
)
#: Ic tablo: isveren sutunlari + ic sutunlar.
INTERNAL_EXTRA_HEADERS: tuple[str, ...] = (
    "Adam-saat",
    "Maliyet B.F.",
    "GG %",
    "Kâr %",
    "Maliyet",
    "GG",
    "Kâr",
)
INTERNAL_HEADERS = EMPLOYER_HEADERS + INTERNAL_EXTRA_HEADERS

LABEL_SUBTOTAL = "Ara Toplam"
LABEL_NET = "NET (KDV Hariç)"
LABEL_GRAND_TOTAL = "GENEL TOPLAM"
LABEL_INTERNAL_TOTALS = "İÇ TOPLAMLAR"


def _col(header: str) -> int:
    """Baslik listesinden 1 tabanli sutun no (ic basliklar isveren basliklarinin devamidir)."""
    return INTERNAL_HEADERS.index(header) + 1


_AMOUNT_COL = _col("Tutar")
_SOURCE_CODE_COL = _col("Bakanlık No")

_INDEX_LABELS = {
    PriceIndexType.ufe: "ÜFE",
    PriceIndexType.tufe: "TÜFE",
    PriceIndexType.construction_cost: "İnşaat Maliyet Endeksi",
    PriceIndexType.fixed_coefficient: "Sabit Katsayı",
}

Cell = str | None
Row = tuple[Cell, ...]


def _s(value: object | None) -> Cell:
    """`None` → BOS hucre (maske / girilmemis); digerleri `str`. Float asla yazilmaz."""
    return None if value is None else str(value)


def _plain(value: Decimal) -> str:
    """Yuzde gosterimi: "20.00" → "20" (girdi yuzdesi; para degil)."""
    return f"{value.normalize():f}"


def _sum(values: Sequence[Decimal | None]) -> Decimal | None:
    """Biri maskeli (`None`) ise toplam da maskeli: kismi toplam YANLIS sayidir."""
    if any(v is None for v in values):
        return None
    return sum((v for v in values if v is not None), Decimal("0.00"))


# ------------------------------------------------------------------ kalem satirlari


def _employer_cells(item: OfferItemRead) -> Row:
    """Fiyatsiz kalemde `customer` `None`: B.F./tutar BOS. Miktar `None` (girilmedi/maske) BOS."""
    customer = item.customer
    return (
        item.poz_no,
        item.source_code,
        item.description,
        item.unit,
        _s(item.quantity),
        _s(customer.unit_price) if customer else None,
        _s(customer.amount) if customer else None,
    )


def _internal_cells(item: OfferItemRead, revision: OfferRevisionRead) -> Row:
    internal = item.internal
    gg_pct = item.overhead_pct if item.overhead_pct is not None else revision.overhead_pct
    return _employer_cells(item) + (
        _s(internal.man_hours),
        _s(item.cost_unit_price),
        _s(gg_pct) if item.priced else None,
        _s(internal.profit_pct) if item.priced else None,
        _s(internal.cost),
        _s(internal.overhead),
        _s(internal.profit),
    )


#: Ara toplam sutunlari: (1 tabanli sutun, kalemden deger, para mi). Para sutunlari YALNIZ
#: toplamlara giren kalemleri (`_enters_totals`) toplar — `calc` ile AYNI kural; adam-saat para
#: degildir: miktarli her kalemin (fiyatsiz dahil) adam-saati toplanir, miktarsizin `None`dir.
_Getter = Callable[[OfferItemRead], Decimal | None]
_SUBTOTALS_EMPLOYER: tuple[tuple[int, _Getter, bool], ...] = (
    (_AMOUNT_COL, lambda i: i.customer.amount if i.customer else None, True),
)
_SUBTOTALS_INTERNAL: tuple[tuple[int, _Getter, bool], ...] = (
    *_SUBTOTALS_EMPLOYER,
    (_col("Adam-saat"), lambda i: i.internal.man_hours, False),
    (_col("Maliyet"), lambda i: i.internal.cost, True),
    (_col("GG"), lambda i: i.internal.overhead, True),
    (_col("Kâr"), lambda i: i.internal.profit, True),
)


def _enters_totals(item: OfferItemRead) -> bool:
    """`calc`in `priced and quantified` kurali, zarftan okunur. Miktar `operasyonel`, tutar `para`
    kovasindadir (ikisi ayri kapsamda maskelenir): tutar VEYA miktar doluysa kalem miktarlidir.
    Ikisi de `None` ise kalem miktarsizdir (tutar uretilmedi) ve ara toplami BOSALTMAZ."""
    if not item.priced:
        return False
    return item.quantity is not None or (
        item.customer is not None and item.customer.amount is not None
    )


def _subtotal_cell(items: Sequence[OfferItemRead], getter: _Getter, *, money: bool) -> Cell:
    """Para: toplamlara giren kalemlerden biri maskeliyse (`None`) BOS (kismi toplam yazilmaz).
    Adam-saat: `None` (miktarsiz) kalemler atlanir; kimlik kovasi oldugu icin maskelenmez."""
    if money:
        values = [getter(i) for i in items if _enters_totals(i)]
    else:
        values = [v for v in (getter(i) for i in items) if v is not None]
    return _s(_sum(values))


def _subtotal_row(items: Sequence[OfferItemRead], view: ExportView, width: int) -> Row:
    cells: list[Cell] = [None] * width
    cells[0] = LABEL_SUBTOTAL
    specs = _SUBTOTALS_EMPLOYER if view is ExportView.employer else _SUBTOTALS_INTERNAL
    for column, getter, money in specs:
        cells[column - 1] = _subtotal_cell(items, getter, money=money)
    return tuple(cells)


# ------------------------------------------------------------------ yazim yardimcilari


def _put(sheet: Worksheet, row: int, values: Sequence[Cell], *, bold: bool = False) -> None:
    for column, value in enumerate(values, start=1):
        if value is None:
            continue  # maskeli/girilmemis: hucreye HIC dokunulmaz (bos)
        if column == _SOURCE_CODE_COL:
            write_text_cell(sheet, row, column, value)  # kod formul olarak calismaz (KAT-B3)
        else:
            sheet.cell(row=row, column=column).value = value
        cell = sheet.cell(row=row, column=column)
        if bold:
            cell.font = Font(bold=True)


def _label_value(sheet: Worksheet, row: int, label: str, value: Cell) -> None:
    _put(sheet, row, (label, value), bold=False)
    sheet.cell(row=row, column=1).font = Font(bold=True)


def _total_row(sheet: Worksheet, row: int, label: str, value: Cell, *, value_col: int) -> None:
    sheet.cell(row=row, column=1).value = label
    sheet.cell(row=row, column=1).font = Font(bold=True)
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=value_col - 1)
    if value is not None:
        sheet.cell(row=row, column=value_col).value = value
        sheet.cell(row=row, column=value_col).font = Font(bold=True)


def _escalation_text(revision: OfferRevisionRead) -> str:
    if revision.price_escalation is OfferPriceEscalation.fixed:
        return "Sabit fiyat"
    index = (
        _INDEX_LABELS[revision.price_index_type] if revision.price_index_type is not None else ""
    )
    return f"Endeksli ({index})" if index else "Endeksli"


# ------------------------------------------------------------------ bolumler


def _write_header_block(
    sheet: Worksheet, offer: Offer, revision: OfferRevisionRead, *, view: ExportView
) -> int:
    title = "TEKLİF" if view is ExportView.employer else "TEKLİF (İÇ)"
    sheet.cell(row=1, column=1).value = title
    sheet.cell(row=1, column=1).font = Font(bold=True, size=14)
    rows: list[tuple[str, Cell]] = [
        ("Teklif No", f"{revision.offer_no} Rev.{revision.rev_no}"),
        ("Tarih", revision.offer_date.strftime(DATE_FORMAT)),
        ("Geçerlilik Bitişi", revision.valid_until.strftime(DATE_FORMAT)),
        ("İşveren", offer.employer_name),
        ("İş Adı", offer.title),
    ]
    if offer.scope_summary:
        rows.append(("Kapsam", offer.scope_summary))
    row = 2
    for label, value in rows:
        _label_value(sheet, row, label, value)
        row += 1
    return row + 1  # bir bos satir


def _write_items(sheet: Worksheet, row: int, revision: OfferRevisionRead, view: ExportView) -> int:
    headers = EMPLOYER_HEADERS if view is ExportView.employer else INTERNAL_HEADERS
    _put(sheet, row, headers, bold=True)
    row += 1
    cells_of: Callable[[OfferItemRead], Row] = (
        _employer_cells
        if view is ExportView.employer
        else (lambda item: _internal_cells(item, revision))
    )
    for index, group in enumerate(revision.groups, start=1):
        sheet.cell(row=row, column=1).value = f"{index}. {group.name}"
        sheet.cell(row=row, column=1).font = Font(bold=True)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        row += 1
        for item in group.items:
            _put(sheet, row, cells_of(item))
            row += 1
        _put(sheet, row, _subtotal_row(group.items, view, len(headers)), bold=True)
        row += 1
    return row + 1


def _write_totals(sheet: Worksheet, row: int, revision: OfferRevisionRead) -> int:
    customer = revision.totals.customer
    _total_row(sheet, row, LABEL_NET, _s(customer.net), value_col=_AMOUNT_COL)
    _total_row(
        sheet,
        row + 1,
        f"KDV (%{_plain(revision.vat_pct)})",
        _s(customer.vat),
        value_col=_AMOUNT_COL,
    )
    _total_row(sheet, row + 2, LABEL_GRAND_TOTAL, _s(customer.gross), value_col=_AMOUNT_COL)
    return row + 4


def _write_internal_totals(sheet: Worksheet, row: int, revision: OfferRevisionRead) -> int:
    totals = revision.totals
    internal = totals.internal
    sheet.cell(row=row, column=1).value = LABEL_INTERNAL_TOTALS
    sheet.cell(row=row, column=1).font = Font(bold=True)
    row += 1
    for label, value in (
        ("Maliyet", internal.cost),
        ("GG", internal.overhead),
        ("Kâr", internal.profit),
        ("Genel Kâr %", internal.profit_pct),
        ("Toplam Adam-saat", internal.man_hours),
        ("Fiyatsız Kalem Sayısı", totals.unpriced_count),
    ):
        _label_value(sheet, row, label, _s(value))
        row += 1
    return row + 1


def _write_conditions(sheet: Worksheet, row: int, revision: OfferRevisionRead) -> int:
    sheet.cell(row=row, column=1).value = "KOŞULLAR"
    sheet.cell(row=row, column=1).font = Font(bold=True)
    row += 1
    conditions: list[tuple[str, Cell]] = [
        ("Ödeme Koşulu", revision.payment_terms),
        (
            "Teslim Süresi",
            f"{revision.delivery_days} gün" if revision.delivery_days is not None else None,
        ),
        ("Fiyat Farkı", _escalation_text(revision)),
        ("Notlar", revision.notes),
    ]
    for label, value in conditions:
        if value:
            _label_value(sheet, row, label, value)
            row += 1
    return row


def _apply_layout(sheet: Worksheet, column_count: int) -> None:
    widths = (14, 16, 42, 10, 14, 16, 16, 12, 16, 10, 10, 16, 16, 16)
    for index in range(column_count):
        sheet.column_dimensions[sheet.cell(row=1, column=index + 1).column_letter].width = widths[
            index
        ]


def build_offer_workbook(offer: Offer, revision: OfferRevisionRead, view: ExportView) -> BytesIO:
    """Revizyon zarfindan xlsx kitabi uretir. Zarf kapsam maskesinden GECMIS olmalidir.

    Bos revizyonda bile gecerli dosya doner (baslik + toplamlar).
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_TITLE
    row = _write_header_block(sheet, offer, revision, view=view)
    row = _write_items(sheet, row, revision, view)
    row = _write_totals(sheet, row, revision)
    if view is ExportView.internal:
        row = _write_internal_totals(sheet, row, revision)
    _write_conditions(sheet, row, revision)
    _apply_layout(sheet, len(EMPLOYER_HEADERS if view is ExportView.employer else INTERNAL_HEADERS))
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
